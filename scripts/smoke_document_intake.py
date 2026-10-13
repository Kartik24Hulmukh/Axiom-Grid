"""Run inside the release container; exercise authenticated real-file intake."""
import json
import os
import tempfile
import urllib.request
from pathlib import Path
from urllib.parse import urlencode

from docx import Document


def write_pdf(path):
    # Minimal valid PDF with compressed content, so decoding bytes is not extraction.
    import zlib
    stream = zlib.compress(b"BT /F1 12 Tf 72 720 Td (CLASSIFICATION: UNCLASSIFIED) Tj 0 -20 Td (SUBJECT: release intake) Tj ET")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" /Filter /FlateDecode >>\nstream\n" + stream + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)
    return path


def main(base_url="http://127.0.0.1:8765", api_key=None):
    """Exercise public owned-upload intake, not a forbidden host-file path."""
    headers = {"Authorization": "Bearer " + (
        os.environ["AXIOM_API_KEYS"] if api_key is None else api_key
    )}
    with tempfile.TemporaryDirectory(prefix="axiom-intake-") as directory:
        root = Path(directory)
        pdf = write_pdf(root / "memo.pdf")
        word = root / "memo.docx"
        doc = Document()
        doc.add_paragraph("CLASSIFICATION: UNCLASSIFIED")
        doc.add_paragraph("SUBJECT: release intake")
        doc.save(word)
        for path in (pdf, word):
            content = path.read_bytes()
            upload = urllib.request.Request(
                base_url + "/api/documents?" + urlencode({"filename": path.name}),
                data=content,
                headers={**headers, "Content-Type": "application/octet-stream"},
            )
            with urllib.request.urlopen(upload, timeout=30) as response:
                owned = json.load(response)
            assert owned["bytes"] == len(content), owned
            assert owned["file"].startswith("uploaded:"), owned
            request = urllib.request.Request(
                base_url + "/api/extract-document",
                data=json.dumps({"file": owned["file"]}).encode(),
                headers={**headers, "Content-Type": "application/json"},
            )
            with urllib.request.urlopen(request, timeout=30) as response:
                data = json.load(response)
            assert data["doc_type"] == "memo", data
            assert data["fields"], data
            assert any(
                "UNCLASSIFIED" in str(field["value"]).upper()
                and field["grounded"] is True
                for field in data["fields"].values()
            ), data
            print(path.suffix, "authenticated real document intake: PASS")


if __name__ == "__main__":
    main()
