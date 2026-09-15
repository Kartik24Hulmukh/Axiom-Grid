"""Exercise real openpyxl ZIP ownership on success and early error return."""
import sys
from pathlib import Path
import pytest
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from sidecar.masters import excel_master

@pytest.mark.parametrize("operations", [[], [{"type": "write_cell", "cell": "not-a-cell", "value": 3}]])
def test_writer_releases_macro_archive(tmp_path, monkeypatch, operations):
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    wb.save(path)
    wb.close()
    loaded = []
    original = excel_master.load_workbook
    def load(*args, **kwargs):
        book = original(*args, **kwargs)
        loaded.append(book)
        return book
    monkeypatch.setattr(excel_master, "load_workbook", load)
    try:
        result = excel_master.ExcelWriter().apply_operations(str(path), operations)
        assert bool(result["errors"]) == bool(operations)
        assert loaded
        archive = loaded[0].vba_archive
        assert archive is None or archive.fp is None
    finally:
        for book in loaded:
            book.close()
            if book.vba_archive is not None:
                book.vba_archive.close()
