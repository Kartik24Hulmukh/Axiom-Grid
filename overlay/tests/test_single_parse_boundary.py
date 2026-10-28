"""Single-parse API boundary regressions.

These are synthetic deterministic extraction fixtures, NOT model/network E2E.
Real parser tests spawn the production child and run the real orchestrator,
request-local SQLite, grounding, provenance and HTTP serialization paths.
"""
from dataclasses import FrozenInstanceError
import hashlib
import pathlib
import threading

import pytest
from fastapi.testclient import TestClient

from kernel.core.data_model import BBox, Chunk, Document, Page
from kernel.sidecar import isolated_ingest
from overlay import server


INVOICE = "Acme Supplies\nInvoice Number: SYN-1042\nTotal Amount: $125.00\nPayment Terms: Net 30\n"


@pytest.fixture(autouse=True)
def synthetic_no_network_gateway(monkeypatch):
    """The deterministic packs need no gateway; any inference is a test error."""
    class ForbiddenGateway:
        def __init__(self, **kwargs):
            pass

        def __getattr__(self, name):
            raise AssertionError(f"Synthetic boundary fixture attempted inference: {name}")

    monkeypatch.setattr(server, "TieredInferenceGateway", ForbiddenGateway)
    monkeypatch.setenv("KAIRO_GATEWAY_TEST_MODE", "true")
    monkeypatch.delenv("AXIOM_EXTRACTION_ISOLATION", raising=False)


def forbidden_parent_parse(*args, **kwargs):
    raise AssertionError("parent IngestorImpl must never parse in process mode")


def graph(path, content, text=INVOICE):
    doc = Document(source_path=str(path), sha256=hashlib.sha256(content).hexdigest(), page_count=1)
    chunks = [Chunk(doc_id=doc.doc_id, page=1, bbox=BBox(.1, .2, .8, .6),
                    text=text, source_type="text", embedding=[.25])]
    return chunks, doc, [Page(doc_id=doc.doc_id, index=1, width_px=800, height_px=1000)]


def source_file(tmp_path, suffix):
    path = tmp_path / ("synthetic-invoice" + suffix)
    if suffix == ".docx":
        from docx import Document as DocxDocument
        source = DocxDocument()
        source.add_paragraph(INVOICE)
        source.save(path)
    elif suffix == ".pdf":
        # Minimal well-formed one-page PDF; no reportlab dependency needed.
        lines = INVOICE.splitlines()
        stream = b"BT /F1 12 Tf 40 750 Td "
        for index, line in enumerate(lines):
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream += (b"0 -20 Td " if index else b"") + f"({escaped}) Tj ".encode("ascii")
        stream += b"ET"
        objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
                   b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                   b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
                   b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
                   f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"]
        data, offsets = b"%PDF-1.4\n", [0]
        for number, obj in enumerate(objects, 1):
            offsets.append(len(data))
            data += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
        xref = len(data)
        data += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
        data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
        data += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
        path.write_bytes(data)
    else:
        path.write_text(INVOICE)
    return path


def observe_real_orchestrator(monkeypatch):
    """Inspect the request store while open; don't substitute a pipeline stub."""
    observed = []
    original = server.OrchestratorImpl.run

    def run(self, doc):
        assert threading.current_thread() is not threading.main_thread()
        assert self._memory is not server.memory_store
        assert self._quality._correction_store is self._memory
        chunks, captured_doc, pages = self._ingestor.ingest(doc.source_path)
        assert captured_doc == doc
        trace = original(self, doc)
        assert self._memory.get_document(doc.doc_id) == doc
        assert self._memory.get_pages(doc.doc_id) == pages
        for chunk in chunks:
            assert self._memory.get_chunk(chunk.chunk_id) == chunk
        assert self._provenance.chunks_for_doc(doc.doc_id) == chunks
        for extraction in trace.extractions:
            chain = self._provenance.get_provenance(extraction.ext_id)
            assert chain.document == doc
            assert chain.chunk.chunk_id == extraction.chunk_id
            if extraction.anchors:
                assert chain.is_complete
                assert extraction.anchors[0].bbox == chain.chunk.bbox
                assert extraction.anchors[0].page == chain.chunk.page
        observed.append((doc, chunks, pages, trace))
        return trace

    monkeypatch.setattr(server.OrchestratorImpl, "run", run)
    return observed


def test_cached_adapter_is_read_only_and_never_reads_source(tmp_path):
    path = tmp_path / "deleted.txt"
    chunks, doc, pages = graph(path, b"original")
    cached = server._CapturedIngestor.capture(chunks, doc, pages)
    chunks[0].embedding.append(9)
    chunks.clear()
    pages.clear()
    first, returned_doc, first_pages = cached.ingest(str(path))
    assert first[0].embedding == [.25]
    first[0].embedding.clear()
    first.clear()
    first_pages.clear()
    second, _, second_pages = cached.ingest(str(path))
    assert len(second) == len(second_pages) == 1
    assert second[0].embedding == [.25]
    assert returned_doc is doc
    with pytest.raises(FrozenInstanceError):
        cached._document = Document()
    with pytest.raises(ValueError, match="different source"):
        cached.ingest(str(path.with_suffix(".pdf")))


@pytest.mark.parametrize("suffix", [".txt", ".md", ".pdf", ".docx"])
def test_real_process_endpoint_preserves_full_graph_without_parent_parser(monkeypatch, tmp_path, suffix):
    path = source_file(tmp_path, suffix)
    content = path.read_bytes()
    calls = []
    original = isolated_ingest.ingest_document_isolated

    def isolated(filepath, timeout, *, content=None):
        assert threading.current_thread() is not threading.main_thread()
        assert 0 < timeout <= server.EXTRACTION_DEADLINE_SECONDS
        calls.append((filepath, content))
        return original(filepath, timeout, content=content)

    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", isolated)
    monkeypatch.setattr(server.IngestorImpl, "ingest", forbidden_parent_parse)
    observed = observe_real_orchestrator(monkeypatch)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 200, response.text
    result = response.json()
    assert calls == [(str(path), content)]
    assert len(observed) == 1
    doc, chunks, pages, trace = observed[0]
    assert doc.source_path == str(path)
    assert doc.sha256 == hashlib.sha256(content).hexdigest()
    assert result["doc_id"] == doc.doc_id
    assert result["doc_type"] == "invoice"
    assert result["fields"]["invoice_number"]["value"] == "SYN-1042"
    assert result["fields"]["total_amount"]["value"] in ("125.00", "$125.00")
    assert doc.page_count == len(pages) == 1
    assert chunks and all(c.doc_id == doc.doc_id and c.bbox is not None and c.page == 1 for c in chunks)
    assert all(p.doc_id == doc.doc_id for p in pages)
    for field in result["fields"].values():
        assert field["grounded"] is True
        assert field["page"] == 1
        assert len(field["bbox"]) == 4
        assert field["source_link"].startswith(f"kairo://doc/{doc.doc_id}?")


def test_captured_bytes_survive_original_source_replacement(monkeypatch, tmp_path):
    path = source_file(tmp_path, ".txt")
    content = path.read_bytes()
    original = isolated_ingest.ingest_document_isolated
    calls = []

    def isolated(filepath, timeout, *, content=None):
        calls.append(content)
        pathlib.Path(filepath).write_text("An unrelated replacement with no invoice signals.")
        return original(filepath, timeout, content=content)

    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", isolated)
    monkeypatch.setattr(server.IngestorImpl, "ingest", forbidden_parent_parse)
    observed = observe_real_orchestrator(monkeypatch)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 200, response.text
    assert calls == [content]
    assert response.json()["doc_type"] == "invoice"
    doc, chunks, _, _ = observed[0]
    assert doc.sha256 == hashlib.sha256(content).hexdigest()
    assert "SYN-1042" in "\n".join(c.text for c in chunks)
    assert "replacement" not in "\n".join(c.text for c in chunks)


def test_explicit_thread_mode_has_one_parser_call_and_snapshot(monkeypatch, tmp_path):
    monkeypatch.setenv("AXIOM_EXTRACTION_ISOLATION", "thread")
    path = source_file(tmp_path, ".docx")
    content = path.read_bytes()
    original = server.IngestorImpl.ingest
    calls = []

    def ingest(self, snapshot):
        assert threading.current_thread() is not threading.main_thread()
        assert snapshot != str(path)
        assert pathlib.Path(snapshot).read_bytes() == content
        calls.append(snapshot)
        path.unlink()  # Classification and orchestration must not reopen it.
        return original(self, snapshot)

    monkeypatch.setattr(server.IngestorImpl, "ingest", ingest)
    observed = observe_real_orchestrator(monkeypatch)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 200, response.text
    assert len(calls) == 1
    assert observed[0][0].source_path == str(path)
    assert observed[0][0].sha256 == hashlib.sha256(content).hexdigest()
    assert response.json()["doc_type"] == "invoice"


@pytest.mark.parametrize("mode", ["", "off", "threads", "PROCESS", " process "])
def test_invalid_isolation_configuration_fails_closed(monkeypatch, tmp_path, mode):
    monkeypatch.setenv("AXIOM_EXTRACTION_ISOLATION", mode)
    path = source_file(tmp_path, ".txt")
    calls = []
    def forbidden(*args, **kwargs):
        calls.append(True)
        raise AssertionError("invalid configuration reached a parser")
    monkeypatch.setattr(server.IngestorImpl, "ingest", forbidden)
    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", forbidden)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 422
    assert calls == []


@pytest.mark.parametrize("suffix", [".xlsx", ".pptx", ".html", ".exe"])
def test_unsupported_type_never_enters_parser(monkeypatch, tmp_path, suffix):
    path = source_file(tmp_path, suffix)
    calls = []
    def forbidden(*args, **kwargs):
        calls.append(True)
        raise AssertionError("unsupported type reached parser")
    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", forbidden)
    monkeypatch.setattr(server.IngestorImpl, "ingest", forbidden)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 422
    assert calls == []


@pytest.mark.parametrize("error,status", [(ValueError("invalid graph"), 422),
    (RuntimeError("missing engine"), 422), (OSError("unavailable child"), 422),
    (isolated_ingest.IsolatedIngestResourceError("resource cap"), 422),
    (isolated_ingest.IsolatedIngestDeadline("timeout"), 504)])
def test_parser_failure_has_no_parent_fallback_and_releases_slot(monkeypatch, tmp_path, error, status):
    path = source_file(tmp_path, ".txt")
    slots = threading.BoundedSemaphore(1)
    calls = []
    parent_calls = []
    def reject(*args, **kwargs):
        calls.append(True)
        raise error
    def parent(*args, **kwargs):
        parent_calls.append(True)
        raise AssertionError("parser failure must not fall back to the parent")
    monkeypatch.setattr(server, "_extraction_slots", slots)
    monkeypatch.setattr(server.IngestorImpl, "ingest", parent)
    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", reject)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == status
    assert calls == [True]
    assert parent_calls == []
    assert slots.acquire(blocking=False)
    slots.release()
    if status == 504:
        assert response.headers["Retry-After"] == "5"


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_real_malformed_document_is_422_not_success_or_parent_fallback(monkeypatch, tmp_path, suffix):
    path = tmp_path / ("invalid" + suffix)
    path.write_bytes(b"Not a valid document")
    monkeypatch.setattr(server.IngestorImpl, "ingest", forbidden_parent_parse)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 422, response.text


def test_real_child_deadline_fails_504_without_parent_parse(monkeypatch, tmp_path):
    path = source_file(tmp_path, ".txt")
    original = isolated_ingest.ingest_document_isolated
    def short_child(filepath, timeout, *, content=None):
        return original(filepath, min(timeout, .001), content=content)
    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", short_child)
    monkeypatch.setattr(server.IngestorImpl, "ingest", forbidden_parent_parse)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 504, response.text
    assert response.headers["Retry-After"] == "5"


@pytest.mark.asyncio
async def test_waiter_deadline_keeps_slot_until_worker_done(monkeypatch, tmp_path):
    import asyncio
    path = source_file(tmp_path, ".txt")
    slots = threading.BoundedSemaphore(1)
    entered, release, finished = threading.Event(), threading.Event(), threading.Event()
    def blocked(filepath, timeout, *, content=None):
        entered.set()
        try:
            assert release.wait(5)
            return graph(filepath, content)
        finally:
            finished.set()
    monkeypatch.setattr(server, "_extraction_slots", slots)
    monkeypatch.setattr(server, "EXTRACTION_DEADLINE_SECONDS", .05)
    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", blocked)
    task = asyncio.create_task(server.extract_document(server.ExtractDocumentRequest(file=str(path))))
    try:
        assert await asyncio.to_thread(entered.wait, 3)
        with pytest.raises(server.HTTPException) as error:
            await task
        assert error.value.status_code == 504
        assert not slots.acquire(blocking=False)
        with pytest.raises(server.HTTPException) as saturated:
            await server.extract_document(server.ExtractDocumentRequest(file=str(path)))
        assert saturated.value.status_code == 503
    finally:
        release.set()
        assert await asyncio.to_thread(finished.wait, 5)
        # Join before checking the done callback, not just the parser's return.
        await asyncio.to_thread(server._extraction_pool.shutdown, wait=True)
    assert slots.acquire(blocking=False)
    slots.release()


@pytest.mark.skipif(__import__("os").name != "posix", reason="POSIX capture protections")
@pytest.mark.parametrize("replacement", ["fifo", "symlink"])
def test_resolved_source_replacement_rejected_without_hanging(tmp_path, replacement):
    """Use a disposable process so regression to blocking open cannot hang CI.

    This checks the final path component only, not an ancestor/path sandbox.
    All input and symlink targets are synthetic test files.
    """
    import os
    import subprocess
    import sys
    import textwrap

    script = textwrap.dedent('''
        import os
        import pathlib
        import sys
        import time
        from fastapi.testclient import TestClient
        from overlay import server
        from kernel.sidecar import isolated_ingest

        root, replacement = pathlib.Path(sys.argv[1]), sys.argv[2]
        path = root / "race.txt"
        path.write_text("Original synthetic regular document.")
        target = root / "target.txt"
        target.write_text("Synthetic symlink target; must never be read.")
        original_resolve = server.resolve_sandbox_path
        resolves, parses = [], []
        def resolve(requested):
            resolved = original_resolve(requested)
            resolves.append(str(resolved))
            resolved.unlink()
            if replacement == "fifo":
                os.mkfifo(resolved)
            else:
                resolved.symlink_to(target)
            return resolved
        def forbidden(*args, **kwargs):
            parses.append(True)
            raise AssertionError("replaced source must never reach parser")
        server.resolve_sandbox_path = resolve
        server.IngestorImpl.ingest = forbidden
        isolated_ingest.ingest_document_isolated = forbidden
        server.EXTRACTION_DEADLINE_SECONDS = .2
        started = time.monotonic()
        with TestClient(server.app) as client:
            response = client.post("/api/extract-document", json={"file": str(path)})
        assert response.status_code == 422, response.text
        assert len(resolves) == 1 and parses == []
        assert time.monotonic() - started < 3
        print("capture race rejected:", replacement)
    ''')
    env = dict(os.environ, AXIOM_STATE_DIR=str(tmp_path / "race-state"), AXIOM_LOG_FORMAT="text")
    completed = subprocess.run([sys.executable, "-c", script, str(tmp_path), replacement],
                               cwd=pathlib.Path(__file__).resolve().parents[2], env=env,
                               capture_output=True, text=True, timeout=10)
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert f"capture race rejected: {replacement}" in completed.stdout


def test_safe_capture_unavailable_fails_closed(monkeypatch, tmp_path):
    path = source_file(tmp_path, ".txt")
    monkeypatch.delattr(server.os, "O_NOFOLLOW")
    calls = []
    def forbidden(*args, **kwargs):
        calls.append(True)
        raise AssertionError("unsafe platform must not enter parser")
    monkeypatch.setattr(isolated_ingest, "ingest_document_isolated", forbidden)
    with TestClient(server.app) as client:
        response = client.post("/api/extract-document", json={"file": str(path)})
    assert response.status_code == 422
    assert calls == []
