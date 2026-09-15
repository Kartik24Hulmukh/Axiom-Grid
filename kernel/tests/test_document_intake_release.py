"""Real file regressions; no model server, renderer or PDF mocks."""
import builtins

import pytest

from kernel.sidecar.ingestor import IngestorImpl


def write_pdf(path):
    # Minimal valid PDF with compressed content, so decoding bytes is not extraction.
    import zlib
    stream = zlib.compress(b"BT /F1 12 Tf 72 720 Td (CLASSIFICATION: UNCLASSIFIED) Tj 0 -20 Td (SUBJECT: release intake) Tj ET")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" /Filter /FlateDecode >>\nstream\n" + stream + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)
    return path


def test_pdf_extracts_real_text_without_optional_model_or_agpl_engine(tmp_path, monkeypatch):
    original = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name.split(".")[0] in {"docling", "fitz"}:
            raise AssertionError("launch ingestion must not invoke optional model/AGPL engines")
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    chunks, doc, pages = IngestorImpl().ingest(str(write_pdf(tmp_path / "memo.pdf")))
    assert "CLASSIFICATION: UNCLASSIFIED" in "\n".join(c.text for c in chunks)
    assert doc.page_count == len(pages) == 1
    assert all(c.page == 1 and 0 <= c.bbox.x0 <= c.bbox.x1 <= 1 and 0 <= c.bbox.y0 <= c.bbox.y1 <= 1 for c in chunks)
    assert all(c.source_type == "pdf_text" for c in chunks)


def test_malformed_pdf_is_not_decoded_as_document_text(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"CLASSIFICATION: UNCLASSIFIED\nNot a PDF")
    with pytest.raises(ValueError, match="PDF"):
        IngestorImpl().ingest(str(path))


def test_missing_docx_parser_fails_closed(tmp_path, monkeypatch):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"this is not extracted Office text")
    original = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("dependency absent")
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    with pytest.raises(RuntimeError, match="python-docx"):
        IngestorImpl().ingest(str(path))


def test_binary_documents_classified_from_parsed_content(tmp_path):
    from docx import Document
    from fastapi.testclient import TestClient
    from overlay.server import app
    pdf = write_pdf(tmp_path / "memo.pdf")
    word = tmp_path / "memo.docx"
    doc = Document()
    doc.add_paragraph("CLASSIFICATION: UNCLASSIFIED")
    doc.add_paragraph("SUBJECT: release intake")
    doc.save(word)
    with TestClient(app) as client:
        for path in (pdf, word):
            result = client.post("/api/extract-document", json={"file": str(path)})
            assert result.status_code == 200, result.text
            assert result.json()["doc_type"] == "memo", result.text
            assert result.json()["fields"], result.text
        bad = tmp_path / "bad.pdf"
        bad.write_bytes(b"Not a PDF")
        assert client.post("/api/extract-document", json={"file": str(bad)}).status_code == 422
