"""Defensive ingestion contract; runnable by stdlib unittest (no pytest needed)."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import multiprocessing
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from kernel.sidecar import isolated_ingest as boundary


def _payload(snapshot):
    return {
        "version": 1, "status": "ok",
        "document": {"doc_id": "document-fixed", "source_path": snapshot,
                     "sha256": hashlib.sha256(Path(snapshot).read_bytes()).hexdigest(),
                     "page_count": 1, "ingested_at": "2026-10-28T15:00:00+00:00"},
        "pages": [{"doc_id": "document-fixed", "index": 1, "width_px": 800,
                   "height_px": 1000, "image_sha256": "a" * 64}],
        "chunks": [{"chunk_id": "chunk-fixed", "doc_id": "document-fixed", "page": 1,
                    "bbox": {"x0": .1, "y0": .2, "x1": .8, "y1": .9},
                    "text": "captured text", "source_type": "docx_paragraph",
                    "embedding": [.25, -.5]}],
    }


# Module-level spawn targets: these are real child processes, not mocked timers.
def _hostile_child(snapshot, result, limits):
    os.setsid()
    mode = os.environ["AXIOM_ASSURANCE_CHILD"]
    if mode == "hang":
        Path(os.environ["AXIOM_ASSURANCE_PID"]).write_text(str(os.getpid()))
        while True:
            time.sleep(1)
    if mode == "missing":
        return
    if mode == "malformed":
        Path(result).write_text("{bad json")
    elif mode == "oversized":
        Path(result).write_bytes(b" " * (limits["output"] + 1))
    elif mode == "fifo":
        os.mkfifo(result)
    elif mode == "symlink":
        os.symlink(snapshot, result)
    elif mode == "hardlink":
        os.link(snapshot, result)
    elif mode == "duplicate":
        Path(result).write_text('{"version":1,"version":1,"status":"reject"}')
    elif mode == "nonfinite":
        Path(result).write_text('{"version":1,"status":"ok","unknown":1e999}')
    elif mode == "deep":
        Path(result).write_text("[" * 2000 + "]" * 2000)
    elif mode == "bad_hash":
        payload = _payload(snapshot)
        payload["document"]["sha256"] = "b" * 64
        boundary._write_result(result, payload, limits["output"])
    elif mode in {"invalid_doc_id", "invalid_chunk_id"}:
        payload = _payload(snapshot)
        if mode == "invalid_doc_id":
            bad_id = "document?injected=1#fragment"
            payload["document"]["doc_id"] = bad_id
            payload["chunks"][0]["doc_id"] = bad_id
            payload["pages"][0]["doc_id"] = bad_id
        else:
            payload["chunks"][0]["chunk_id"] = "chunk\x00injected"
        boundary._write_result(result, payload, limits["output"])
    elif mode == "roundtrip":
        boundary._write_result(result, _payload(snapshot), limits["output"])
    elif mode == "mutate_original":
        Path(os.environ["AXIOM_ASSURANCE_ORIGINAL"]).write_text("new original bytes")
        payload = _payload(snapshot)
        payload["chunks"][0]["text"] = Path(snapshot).read_text()
        boundary._write_result(result, payload, limits["output"])
    else:
        os._exit(9)


def _limits_failure_child(snapshot, result, limits):
    import resource
    # Raising a real hard limit is not reliable under a privileged test user.
    # Inject failure inside a real spawned process, before any parser import.
    with patch.object(resource, "setrlimit", side_effect=OSError("unsupported limit")):
        _REAL_CHILD_ENTRY(snapshot, result, limits)


_REAL_CHILD_ENTRY = boundary._child_entry


@unittest.skipUnless(os.name == "posix", "Boundary deliberately fails closed outside POSIX")
class IsolatedIngestionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name)
        self.path = self.directory / "original.txt"
        self.path.write_bytes(b"first paragraph\n\nsecond paragraph")
        settings = {k: v for k, v in os.environ.items()
                    if not (k.startswith("AXIOM_INGEST_") or k.startswith("AXIOM_ASSURANCE_")
                            or k in {"AXIOM_EXTRACTION_ISOLATION", "AXIOM_MAX_EXTRACTED_CHARS"})}
        self.env = patch.dict(os.environ, settings, clear=True)
        self.env.start()
        self.addCleanup(self.env.stop)

    def ingest(self, **kwargs):
        return boundary.ingest_document_isolated(str(self.path), 20, **kwargs)

    def hostile(self, mode, **kwargs):
        with patch.dict(os.environ, {"AXIOM_ASSURANCE_CHILD": mode}), patch.object(
                boundary, "_child_entry", _hostile_child):
            return self.ingest(**kwargs)

    def test_real_text_full_graph_and_compatibility(self):
        self.assertTrue(boundary.isolation_enabled())
        chunks, doc, pages = self.ingest()
        self.assertEqual([c.text for c in chunks], ["first paragraph", "second paragraph"])
        self.assertEqual(doc.sha256, hashlib.sha256(self.path.read_bytes()).hexdigest())
        self.assertEqual(doc.source_path, str(self.path))
        self.assertIsNotNone(doc.ingested_at.utcoffset())
        self.assertEqual(len(pages), doc.page_count)
        for chunk in chunks:
            self.assertEqual(chunk.doc_id, doc.doc_id)
            self.assertIn(chunk.page, {p.index for p in pages})
            self.assertIsNotNone(chunk.bbox)
            self.assertEqual(chunk.source_type, "text")
        self.assertEqual(boundary.ingest_isolated(str(self.path), 20), [c.text for c in chunks])

    def test_real_docx_subprocess_from_captured_bytes_only(self):
        from docx import Document
        source = Document()
        source.add_paragraph("Real Office paragraph")
        source.add_paragraph("Provenance survives")
        buf = io.BytesIO()
        source.save(buf)
        content = buf.getvalue()
        original = self.directory / "never-created.DOCX"
        chunks, doc, pages = boundary.ingest_document_isolated(str(original), 20, content=content)
        self.assertFalse(original.exists())
        self.assertEqual([c.text for c in chunks], ["Real Office paragraph", "Provenance survives"])
        self.assertEqual(doc.sha256, hashlib.sha256(content).hexdigest())
        self.assertEqual(doc.source_path, str(original))
        self.assertTrue(all(c.source_type == "docx_paragraph" for c in chunks))
        self.assertTrue(all(p.doc_id == doc.doc_id for p in pages))

    def test_capture_wins_over_mutated_original(self):
        os.environ["AXIOM_ASSURANCE_ORIGINAL"] = str(self.path)
        captured = b"immutable capture"
        chunks, doc, _ = self.hostile("mutate_original", content=captured)
        self.assertEqual(self.path.read_text(), "new original bytes")
        self.assertEqual(chunks[0].text, captured.decode())
        self.assertEqual(doc.sha256, hashlib.sha256(captured).hexdigest())
        self.assertEqual(doc.source_path, str(self.path))

    def test_metadata_ids_geometry_embeddings_round_trip(self):
        chunks, doc, pages = self.hostile("roundtrip")
        self.assertEqual(doc.doc_id, "document-fixed")
        self.assertEqual(chunks[0].chunk_id, "chunk-fixed")
        self.assertEqual(chunks[0].embedding, [.25, -.5])
        self.assertEqual(chunks[0].bbox.x0, .1)
        self.assertEqual(chunks[0].bbox.y1, .9)
        self.assertEqual(pages[0].image_sha256, "a" * 64)
        self.assertEqual((pages[0].width_px, pages[0].height_px), (800, 1000))
        self.assertEqual(doc.ingested_at.isoformat(), "2026-10-28T15:00:00+00:00")

    def test_route_identifiers_reject_child_punctuation_and_controls(self):
        for mode in ("invalid_doc_id", "invalid_chunk_id"):
            with self.subTest(mode=mode), self.assertRaises(ValueError):
                self.hostile(mode)
        for bad in ("a?query=1", "a#fragment", "a/b", "a\\b", "a%00", "a\x00b", "a\nb",
                    "a b", "a.b", "é", "", "a" * 257):
            with self.subTest(identifier=repr(bad)), self.assertRaises(ValueError):
                boundary._identifier(bad)
        for good in ("document-fixed", "chunk-fixed", "UPPER_lower-0123",
                     "550e8400-e29b-41d4-a716-446655440000", "a" * 256):
            self.assertEqual(boundary._identifier(good), good)

    def test_child_output_is_untrusted(self):
        os.environ["AXIOM_INGEST_MAX_OUTPUT_BYTES"] = "8192"
        for mode in ("malformed", "missing", "oversized", "fifo", "symlink", "hardlink",
                     "duplicate", "deep", "nonfinite", "bad_hash", "exit"):
            with self.subTest(mode=mode):
                with self.assertRaises((ValueError, boundary.IsolatedIngestResourceError)):
                    self.hostile(mode)

    def test_real_large_response_has_no_feeder_deadlock(self):
        content = b"a" * (1024 * 1024)
        started = time.monotonic()
        chunks, doc, _ = self.ingest(content=content)
        self.assertEqual(chunks[0].text, content.decode())
        self.assertEqual(doc.sha256, hashlib.sha256(content).hexdigest())
        self.assertLess(time.monotonic() - started, 20)

    def test_actual_hung_process_is_killed_reaped_and_private_files_removed(self):
        marker = self.directory / "pid"
        os.environ["AXIOM_ASSURANCE_PID"] = str(marker)
        before = {p.pid for p in multiprocessing.active_children()}
        started = time.monotonic()
        with patch.dict(os.environ, {"AXIOM_ASSURANCE_CHILD": "hang"}), patch.object(
                boundary, "_child_entry", _hostile_child), patch.object(
                boundary.tempfile, "tempdir", str(self.directory)):
            with self.assertRaises(boundary.IsolatedIngestDeadline):
                boundary.ingest_document_isolated(str(self.path), 1.5)
        self.assertLess(time.monotonic() - started, 5)
        self.assertTrue(marker.exists(), "Child must really start and hang")
        pid = int(marker.read_text())
        with self.assertRaises(ProcessLookupError):
            os.kill(pid, 0)
        self.assertEqual({p.pid for p in multiprocessing.active_children()}, before)
        self.assertEqual(list(self.directory.glob("axiom-ingest-*")), [])

    def test_invalid_deadlines(self):
        for value in (0, -1, float("nan"), float("inf"), -float("inf"), True, "1", None, 10**1000):
            with self.subTest(value=str(value)[:40]), self.assertRaises(ValueError):
                boundary.ingest_document_isolated(str(self.path), value)

    def test_configuration_is_strict_and_process_default(self):
        for bad in ("", "Process", " process", "false", "threads"):
            with patch.dict(os.environ, {"AXIOM_EXTRACTION_ISOLATION": bad}), self.assertRaises(ValueError):
                self.ingest()
        for name in boundary._configuration():
            self.assertGreater(boundary._configuration()[name], 0)
        for name in ("AXIOM_INGEST_CPU_SECONDS", "AXIOM_INGEST_RSS_BYTES", "AXIOM_INGEST_MAX_INPUT_BYTES",
                     "AXIOM_INGEST_MAX_OUTPUT_BYTES", "AXIOM_INGEST_MAX_CHUNKS", "AXIOM_INGEST_MAX_PAGES",
                     "AXIOM_MAX_EXTRACTED_CHARS"):
            for bad in ("0", "-1", "nan", "inf", "2.5", " 2", "9999999999999999999999"):
                with self.subTest(name=name, value=bad), patch.dict(os.environ, {name: bad}):
                    with self.assertRaises(ValueError):
                        self.ingest()

    def test_limits_failure_does_not_run_parser(self):
        with patch.object(boundary, "_child_entry", _limits_failure_child):
            with self.assertRaises(boundary.IsolatedIngestResourceError):
                self.ingest()

    def test_explicit_thread_dev_mode_still_validates(self):
        with patch.dict(os.environ, {"AXIOM_EXTRACTION_ISOLATION": "thread"}):
            self.assertFalse(boundary.isolation_enabled())
            chunks, doc, _ = self.ingest(content=b"dev text")
        self.assertEqual(chunks[0].text, "dev text")
        self.assertEqual(doc.sha256, hashlib.sha256(b"dev text").hexdigest())

    def test_input_caps_types_suffix_and_regular_file(self):
        with patch.dict(os.environ, {"AXIOM_INGEST_MAX_INPUT_BYTES": "3"}):
            with self.assertRaises(ValueError):
                self.ingest()
            with self.assertRaises(ValueError):
                self.ingest(content=b"four")
        for content in (bytearray(b"a"), memoryview(b"a"), "a"):
            with self.subTest(content=type(content)), self.assertRaises(ValueError):
                self.ingest(content=content)
        fifo = self.directory / "input.txt"
        os.mkfifo(fifo)
        with self.assertRaises(ValueError):
            boundary.ingest_document_isolated(str(fifo), 10)
        for path in ("bad.bin", "", "null\x00.txt"):
            with self.assertRaises(ValueError):
                boundary.ingest_document_isolated(path, 10, content=b"test")

    def test_unavailable_platform_primitives_and_rlimits_fail_closed(self):
        import resource
        with patch.object(boundary.os, "name", "nt"):
            with self.assertRaises(boundary.IsolatedIngestResourceError):
                self.ingest()
        for owner, names in ((boundary.os, ("O_NOFOLLOW", "O_NONBLOCK", "setsid", "killpg")),
                             (resource, ("RLIMIT_CPU", "RLIMIT_AS", "RLIMIT_FSIZE", "RLIMIT_CORE"))):
            for name in names:
                with self.subTest(name=name), patch.object(owner, name, create=True):
                    delattr(owner, name)
                    with self.assertRaises(boundary.IsolatedIngestResourceError):
                        self.ingest()

    def test_direct_input_symlink_is_not_followed(self):
        link = self.directory / "link.txt"
        link.symlink_to(self.path)
        with self.assertRaises(OSError):
            boundary.ingest_document_isolated(str(link), 10)

    def test_rejects_invalid_real_docx(self):
        with self.assertRaises((ValueError, boundary.IsolatedIngestResourceError)):
            boundary.ingest_document_isolated("invalid.docx", 20, content=b"not a zip")

    def test_output_and_graph_count_text_caps(self):
        for settings, content in (({"AXIOM_INGEST_MAX_OUTPUT_BYTES": "512"}, b"x" * 2048),
                                  ({"AXIOM_INGEST_MAX_CHUNKS": "1"}, b"a\n\nb"),
                                  ({"AXIOM_INGEST_MAX_PAGES": "1"}, b"\n".join([b"x"] * 61) + b"\n\ny"),
                                  ({"AXIOM_MAX_EXTRACTED_CHARS": "3"}, b"abcd"),
                                  ({"AXIOM_MAX_EXTRACTED_CHARS": "3"}, b"ab\n\ncd")):
            with self.subTest(settings=settings), patch.dict(os.environ, settings):
                with self.assertRaises((ValueError, boundary.IsolatedIngestResourceError)):
                    self.ingest(content=content)

    def test_schema_types_references_and_hash_consistency(self):
        snapshot = str(self.path)
        payload = _payload(snapshot)
        digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
        limits = boundary._configuration()
        edits = [
            ((), None), (("version",), True), (("status",), "other"),
            (("document", "source_path"), "/wrong"), (("document", "sha256"), "c" * 64),
            (("document", "sha256"), "z" * 64), (("document", "doc_id"), ""),
            (("document", "page_count"), True), (("document", "page_count"), 0),
            (("document", "ingested_at"), "2026-01-01"), (("document", "ingested_at"), 9),
            (("pages",), []), (("pages", 0, "index"), 2), (("pages", 0, "doc_id"), "wrong"),
            (("pages", 0, "width_px"), False), (("pages", 0, "height_px"), 0),
            (("pages", 0, "image_sha256"), "bad"), (("chunks",), {}),
            (("chunks", 0, "doc_id"), "wrong"), (("chunks", 0, "chunk_id"), []),
            (("chunks", 0, "page"), 0), (("chunks", 0, "page"), True),
            (("chunks", 0, "text"), 1), (("chunks", 0, "source_type"), None),
            (("chunks", 0, "bbox"), None), (("chunks", 0, "bbox", "x0"), float("nan")),
            (("chunks", 0, "bbox", "x0"), float("inf")), (("chunks", 0, "bbox", "x0"), True),
            (("chunks", 0, "bbox", "x0"), -.1), (("chunks", 0, "bbox", "x0"), .99),
            (("chunks", 0, "bbox", "x0"), 10**1000), (("chunks", 0, "embedding"), [False]),
            (("chunks", 0, "embedding"), [math.inf]), (("chunks", 0, "embedding"), "bad"),
            (("chunks", 0, "embedding"), [0] * (boundary.MAX_EMBEDDING_VALUES + 1)),
        ]
        for keys, value in edits:
            modified = copy.deepcopy(payload)
            if keys:
                at = modified
                for key in keys[:-1]:
                    at = at[key]
                at[keys[-1]] = value
            else:
                modified = value
            with self.subTest(keys=keys, value=str(value)[:30]), self.assertRaises(ValueError):
                boundary._validate(modified, snapshot, digest, limits)
        modified = copy.deepcopy(payload)
        modified["chunks"].append(copy.deepcopy(modified["chunks"][0]))
        with self.assertRaises(ValueError):
            boundary._validate(modified, snapshot, digest, limits)
        modified = copy.deepcopy(payload)
        modified["document"]["page_count"] = 2
        modified["pages"].append(copy.deepcopy(modified["pages"][0]))
        with self.assertRaises(ValueError):
            boundary._validate(modified, snapshot, digest, limits)
        modified = copy.deepcopy(payload)
        modified["document"]["extra"] = "unexpected"
        with self.assertRaises(ValueError):
            boundary._validate(modified, snapshot, digest, limits)


if __name__ == "__main__":
    unittest.main()
